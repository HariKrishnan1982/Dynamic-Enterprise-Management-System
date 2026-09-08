"""Document ingestion service.

Parses uploaded files or URLs into plain text, chunks them, and writes
KnowledgeChunk rows to the DB. Sets ingestion_status on the source.

Parsers used:
  PDF     → pypdf
  DOCX    → python-docx
  XLSX    → openpyxl
  TXT     → built-in
  URL     → httpx + BeautifulSoup4
  other   → raw bytes decode attempt
"""
import hashlib
import json
import logging
import os
import re
import tempfile
from pathlib import Path
from typing import Generator

logger = logging.getLogger(__name__)

# ─── Chunker ─────────────────────────────────────────────────────────────────

CHUNK_SIZE = int(os.getenv("INGEST_CHUNK_SIZE", "1500"))   # chars
CHUNK_OVERLAP = int(os.getenv("INGEST_CHUNK_OVERLAP", "200"))


def _chunk_text(text: str) -> list[str]:
    """Split text into overlapping chunks on paragraph boundaries."""
    if not text.strip():
        return []

    # Split on blank lines first (paragraph chunks)
    paragraphs: list[str] = re.split(r"\n{2,}", text)
    chunks: list[str] = []
    buffer = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(buffer) + len(para) + 2 <= CHUNK_SIZE:
            buffer = (buffer + "\n\n" + para).lstrip()
        else:
            if buffer:
                chunks.append(buffer)
                # start new buffer with overlap
                overlap_text = buffer[-CHUNK_OVERLAP:] if CHUNK_OVERLAP > 0 else ""
                buffer = (overlap_text + "\n\n" + para).lstrip()
            else:
                # Single paragraph longer than chunk_size — split by sentence
                for sentence in re.split(r"(?<=[.!?])\s+", para):
                    if len(buffer) + len(sentence) + 1 <= CHUNK_SIZE:
                        buffer = (buffer + " " + sentence).lstrip()
                    else:
                        if buffer:
                            chunks.append(buffer)
                        buffer = sentence
    if buffer.strip():
        chunks.append(buffer.strip())
    return chunks


# ─── Parsers ─────────────────────────────────────────────────────────────────

def _parse_pdf(file_path: str) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(p for p in pages if p.strip())
        if not text.strip():
            logger.warning("PDF produced no text (possibly scanned): %s", file_path)
        return text
    except Exception as exc:
        logger.error("PDF parse error for %s: %s", file_path, exc)
        raise


def _parse_docx(file_path: str) -> str:
    try:
        from docx import Document  # type: ignore
        doc = Document(file_path)
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as exc:
        logger.error("DOCX parse error for %s: %s", file_path, exc)
        raise


def _parse_xlsx(file_path: str) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
        wb = load_workbook(file_path, read_only=True, data_only=True)
        parts: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append("\t".join(cells))
            if rows:
                parts.append(f"## Sheet: {sheet_name}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    except Exception as exc:
        logger.error("XLSX parse error for %s: %s", file_path, exc)
        raise


def _parse_txt(file_path: str) -> str:
    try:
        with open(file_path, encoding="utf-8", errors="replace") as f:
            return f.read()
    except Exception as exc:
        logger.error("TXT parse error for %s: %s", file_path, exc)
        raise


def _parse_url(url: str) -> str:
    try:
        import httpx  # type: ignore
        from bs4 import BeautifulSoup  # type: ignore

        headers = {"User-Agent": "ConAI-Crawler/1.0 (+https://conai.example.com/bot)"}
        resp = httpx.get(url, headers=headers, follow_redirects=True, timeout=30)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, "html.parser")
        # Strip boilerplate
        for tag in soup(["script", "style", "noscript", "iframe", "nav", "header", "footer", "aside"]):
            tag.decompose()
        # Prefer <main> or <article>
        content = soup.find("main") or soup.find("article") or soup.find("body") or soup
        text = content.get_text(separator="\n")
        # Collapse whitespace
        lines = [line.strip() for line in text.splitlines()]
        return "\n".join(line for line in lines if line)
    except Exception as exc:
        logger.error("URL parse error for %s: %s", url, exc)
        raise


_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "txt": _parse_txt,
    "url": _parse_url,
}


def extract_text(source_type: str, file_path: str | None, source_url: str | None) -> str:
    """Dispatch to the correct parser and return plain text."""
    if source_type == "url" and source_url:
        return _parse_url(source_url)
    if file_path:
        parser = _PARSERS.get(source_type, _parse_txt)
        return parser(file_path)
    raise ValueError(f"No file_path or source_url for source_type={source_type!r}")


# ─── Background ingestion job ─────────────────────────────────────────────────

def ingest_source(source_id: int, db_factory) -> None:
    """Run in a background thread: parse → chunk → persist chunks.

    db_factory is called to open a *new* session (the caller's session may be
    closed by the time the thread runs).
    """
    db = db_factory()
    try:
        from app.models import KnowledgeChunk, KnowledgeSource  # avoid circular at module level

        source: KnowledgeSource | None = db.get(KnowledgeSource, source_id)
        if not source:
            logger.error("ingest_source: source %d not found", source_id)
            return

        source.ingestion_status = "processing"
        db.commit()

        try:
            text = extract_text(source.source_type, source.file_path, source.source_url)
            chunks = _chunk_text(text)

            # Delete old chunks before re-indexing
            db.query(KnowledgeChunk).filter(KnowledgeChunk.source_id == source_id).delete()

            for idx, chunk_text in enumerate(chunks):
                checksum = hashlib.sha256(chunk_text.encode()).hexdigest()
                chunk = KnowledgeChunk(
                    source_id=source_id,
                    chunk_index=idx,
                    text=chunk_text,
                    metadata_json=json.dumps(
                        {"source_name": source.name, "chunk_index": idx, "checksum": checksum}
                    ),
                )
                db.add(chunk)

            source.ingestion_status = "indexed"
            source.ingestion_error = None
            db.commit()
            logger.info("ingest_source: source %d indexed, %d chunks", source_id, len(chunks))

        except Exception as exc:
            db.rollback()
            source = db.get(KnowledgeSource, source_id)
            if source:
                source.ingestion_status = "failed"
                source.ingestion_error = str(exc)
                db.commit()
            logger.error("ingest_source: source %d failed: %s", source_id, exc)
    finally:
        db.close()
